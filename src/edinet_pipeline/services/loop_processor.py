import os
from datetime import datetime
from time import perf_counter

from openpyxl import load_workbook

from edinet_pipeline.domain.output_buffer import OutputBuffer
from edinet_pipeline.domain.raw_builder import RAW_COLS
from edinet_pipeline.domain.skip import SkipCode, add_skip
from edinet_pipeline.services.loop_stage_service import (
    build_excel_not_found_result,
    build_excel_write_inputs_stage,
    build_raw_rows_stage,
    build_stock_write_context,
    close_workbook_quietly,
    create_loop_event,
    execute_stock_write_stage,
    finalize_output_excel,
    open_workbook_stage,
    pick_company_name,
    pick_period_end,
    run_parse_stages,
    save_workbook_stage,
    write_named_range_stage,
    write_raw_sheet_stage,
)
from edinet_pipeline.services.raw_service import build_raw_rows_all_docs
from edinet_pipeline.services.stock_service import write_stock_data_to_workbook
from edinet_pipeline.services.summary_service import write_loop_summary
from edinet_pipeline.services.template_contract_service import OPTIONAL_TEMPLATE_OUTPUT_NAMES
from edinet_pipeline.services.workbook_service import prepare_workbook


def process_one_loop(loop, date_pairs, skipped_files, logger, parse_cache=None, runtime=None):
    company_code_from_job = loop.get("company_code")
    company_name_from_job = loop.get("company_name")
    has_half_from_job = loop.get("has_half")
    source_zips = loop.get("source_zips") or []
    output_root = loop.get("output_root")

    parsed_docs = []
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    t0 = perf_counter()

    loop_event = create_loop_event(
        loop=loop,
        company_code=company_code_from_job,
        company_name=company_name_from_job,
        has_half=has_half_from_job,
        source_zips=source_zips,
        run_id=run_id,
    )

    out_buffer = OutputBuffer()

    selected_file, excel_file_path, excel_base_name = prepare_workbook(loop, run_id, logger)
    if not selected_file:
        add_skip(
            skipped_files,
            code=SkipCode.EXCEL_NOT_FOUND,
            phase="excel_select",
            loop=loop,
            excel=excel_base_name,
            xbrl=None,
            message="使用するExcelが見つからない",
        )
        logger.warning(
            f"[company failed] slot={loop.get('slot')} "
            f"code={company_code_from_job} "
            f"name={company_name_from_job} "
            f"phase=excel_select"
        )
        return build_excel_not_found_result(
            slot=loop.get("slot"),
            company_code=company_code_from_job,
            company_name=company_name_from_job,
        )

    loop_event["excel"] = os.path.basename(excel_file_path)
    xbrl_file_paths = loop["xbrl_file_paths"]

    logger.info(
        f"[company detect] slot={loop.get('slot')} "
        f"code={company_code_from_job} "
        f"name={company_name_from_job} "
        f"half={has_half_from_job} "
        f"file1={len(xbrl_file_paths.get('file1') or [])} "
        f"file2={len(xbrl_file_paths.get('file2') or [])} "
        f"file3={len(xbrl_file_paths.get('file3') or [])}"
    )

    parse_stage = run_parse_stages(
        loop=loop,
        xbrl_file_paths=xbrl_file_paths,
        excel_file_path=excel_file_path,
        parsed_docs=parsed_docs,
        skipped_files=skipped_files,
        loop_event=loop_event,
        out_buffer=out_buffer,
        logger=logger,
        perf_counter=perf_counter,
        parse_cache=parse_cache,
    )
    x1 = parse_stage["x1"]
    x2 = parse_stage["x2"]
    meta2 = parse_stage["meta2"]
    security_code = parse_stage["security_code"]
    use_half = parse_stage["use_half"]

    out_buffer_dict, display_unit = build_excel_write_inputs_stage(
        out_buffer=out_buffer,
        xbrl_file_paths=xbrl_file_paths,
        x1=x1,
        x2=x2,
        use_half=use_half,
        loop=loop,
        company_code=company_code_from_job,
        security_code=security_code,
        company_name=company_name_from_job,
        parse_cache=parse_cache,
        logger=logger,
    )

    raw_rows = build_raw_rows_stage(
        parsed_docs=parsed_docs,
        security_code=security_code,
        run_id=run_id,
        loop_event=loop_event,
        loop=loop,
        company_code=company_code_from_job,
        logger=logger,
        perf_counter=perf_counter,
        build_raw_rows_func=build_raw_rows_all_docs,
    )
    wb = open_workbook_stage(
        excel_file_path=excel_file_path,
        loop_event=loop_event,
        loop=loop,
        company_code=company_code_from_job,
        security_code=security_code,
        logger=logger,
        perf_counter=perf_counter,
        load_workbook_func=load_workbook,
    )

    stock_date_pairs = []
    stock_status = None
    write_raw_sheet = True if runtime is None else bool(getattr(runtime, "write_raw_sheet", True))
    enable_stock = True if runtime is None else bool(getattr(runtime, "enable_stock", True))

    try:
        write_named_range_stage(
            workbook=wb,
            out_buffer_dict=out_buffer_dict,
            display_unit=display_unit,
            loop_event=loop_event,
            loop=loop,
            company_code=company_code_from_job,
            security_code=security_code,
            company_name=company_name_from_job,
            logger=logger,
            perf_counter=perf_counter,
            optional_output_names=OPTIONAL_TEMPLATE_OUTPUT_NAMES,
        )

        write_raw_sheet_stage(
            workbook=wb,
            raw_rows=raw_rows,
            raw_cols=RAW_COLS,
            write_raw_sheet=write_raw_sheet,
            loop_event=loop_event,
            loop=loop,
            company_code=company_code_from_job,
            security_code=security_code,
            logger=logger,
            perf_counter=perf_counter,
        )

        stock_context = build_stock_write_context(
            out_buffer_dict=out_buffer_dict,
            x1=x1,
            use_half=use_half,
            security_code=security_code,
        )
        stock_date_pairs = stock_context["stock_date_pairs"]
        stock_code = stock_context["stock_code"]
        stock_status = execute_stock_write_stage(
            workbook=wb,
            stock_code=stock_code,
            stock_date_pairs=stock_date_pairs,
            enable_stock=enable_stock,
            loop_event=loop_event,
            loop=loop,
            company_code=company_code_from_job,
            security_code=security_code,
            company_name=company_name_from_job,
            logger=logger,
            perf_counter=perf_counter,
            write_stock_func=write_stock_data_to_workbook,
        )

        save_workbook_stage(
            workbook=wb,
            excel_file_path=excel_file_path,
            loop_event=loop_event,
            loop=loop,
            company_code=company_code_from_job,
            security_code=security_code,
            logger=logger,
            perf_counter=perf_counter,
        )
    finally:
        close_workbook_quietly(wb)
        wb = None

    period_end_date = pick_period_end(x1, x2, meta2)
    final_security_code = security_code or company_code_from_job or ""
    final_company_name = pick_company_name(x1, x2, meta2, company_name_from_job)

    final_excel_file_path = finalize_output_excel(
        excel_file_path=excel_file_path,
        output_root=output_root,
        security_code=final_security_code,
        company_name=final_company_name,
        period_end_date=period_end_date,
        logger=logger,
    )

    loop["final_excel_file_path"] = final_excel_file_path
    loop_event["excel"] = os.path.basename(final_excel_file_path)
    loop_event["company_name"] = final_company_name

    write_loop_summary(
        loop_event=loop_event,
        security_code=final_security_code,
        raw_rows=raw_rows,
        out_buffer_dict=out_buffer_dict,
        skipped_files=skipped_files,
        loop=loop,
        t0=t0,
        perf_counter=perf_counter,
        logger=logger,
    )

    logger.info(
        f"[company done] slot={loop.get('slot')} "
        f"code={final_security_code} "
        f"name={final_company_name} "
        f"mode={'half' if use_half else 'full'} "
        f"output_excel={final_excel_file_path}"
    )

    return {
        "slot": loop.get("slot"),
        "company_code": final_security_code,
        "company_name": final_company_name,
        "status": "success",
        "stock_status": stock_status,
        "output_excel": final_excel_file_path,
    }
